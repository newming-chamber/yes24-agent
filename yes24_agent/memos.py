"""Yes24 제공 메모 — 사업부가 전해 준 "누가 무엇을 말했는가"를 메모 한 건씩 담는다.

선정위원 추천사·작가의 말·캠페인 일정처럼 사이트에 없는 1차 자료가 대상이다(설계서
design/memo-model.md). 메모 한 건 = 내용 + 누가(선택) + 링크(선택, Yes24 주소 하나) + 관련
상품번호(0개 이상)이고,
켜기/끄기·기간·설명(언제 쓰는 자료인지)은 메모를 담는 **묶음**에 둔다. 표·열 스키마는 없다.

- 가져오기(`read_import`): 엑셀/CSV는 메모를 한꺼번에 만드는 보조 입력이다. 첫 비어 있지 않은
  행이 머리글이고, 사람이 고른 내용 열들을 "열이름: 값" 줄로 이어 붙여 메모 내용을 만든다(열이
  하나면 값만). 시트
  셀은 스트리밍으로 읽고(행·열 상한), 공유 문자열처럼 통째로 올라가는 부분은 압축 해제 크기 상한이
  묶는다. 어드민과 CLI가 같은 함수를 지난다. 개인 연락처
  형태(이메일·전화)는 저장 전에 거부한다 — 답변 필터가 아니라 받지 않을 자료를 막는 보안 검증이다.
- 저장(`MemoService`): 세션 DB 안의 테이블 두 개(묶음·메모). 모든 쓰기의 첫 문장은 묶음 행의
  조건부 UPDATE(revision+1)라, 같은 묶음의 쓰기는 그 행 잠금으로 한 줄로 서고 낡은 화면은 409를
  받는다. 버전·롤백은 없다 — 사고 대응은 끄기이고, 이력은 같은 트랜잭션의 감사 기록뿐이다. 캐시도
  없다 — 켜기·끄기·수정이 다음 조회부터 모든 서버에 반영된다.
- 투영(`memo_view`·`search_view`): 순위 없는 부분 일치, 묶음 id → 메모 id 순. 조회 도구·CLI
  `show`가 같은 함수를 쓴다.

가격·재고·판매상태는 메모로 말하지 않는다(4a) — 메모 출처의 url은 링크·묶음 관련 페이지·첫 상품·
홈 중 앞선 것에 `#notice-{id}`를 붙인 것이라 같은 턴 상품 출처와 병합되지 않는다.
"""

from __future__ import annotations

import contextlib
import csv
import hashlib
import io
import logging
import re
import unicodedata
import zipfile
from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import asdict, dataclass, field, fields
from datetime import date, datetime, time, timezone
from pathlib import PurePath
from typing import Any
from urllib.parse import urlsplit

from fastapi import HTTPException
from sqlalchemy import (
    Boolean,
    Column,
    Date,
    DateTime,
    Index,
    Integer,
    MetaData,
    String,
    Table,
    Text,
    delete,
    func,
    insert,
    select,
    update,
)
from sqlalchemy.dialects import mysql
from sqlalchemy.engine import make_url
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncEngine, create_async_engine
from sqlalchemy.schema import CreateIndex, CreateTable

from yes24_agent.admin_auth import EditConflict
from yes24_agent.config import Settings, get_settings
from yes24_agent.sources import DETAIL_FIDELITY, KST
from yes24_agent.yes24.client import allowed_domain, is_allowed_host
from yes24_agent.yes24.urls import product_url

logger = logging.getLogger(__name__)

MEMO_SOURCE_TYPE = "knowledge"  # 출처 타입(공개 notice로 접힌다 — toolsets)

# ── 스키마 ─────────────────────────────────────────────────────────────────
# DDL의 유일한 원본이다(sqlite·MySQL 공용). 실행은 사람이 CLI `init`으로만 하고 앱 기동 시 자동
# 생성은 없다 — 운영 DDL은 `init --print`로 먼저 본다. FK·CASCADE는 두지 않는다(sqlite는 PRAGMA
# 의존) — 묶음 삭제·통째 교체가 같은 트랜잭션에서 메모를 직접 지운다.
_METADATA = MetaData()
# 길이 상수는 아래 String(n)과 같은 값이다 — 검증이 DDL을 따른다(MySQL strict는 초과 시 DataError).
_TITLE_MAX = 100
_SHORT_MAX = 100  # 누가
_SOURCE_URL_MAX = 500
# 호스트는 도메인 문자만 — 파이썬은 "evil.com\\.yes24.com"을 yes24 하위 도메인으로 읽지만 브라우저는
# 역슬래시를 슬래시로 바꿔 evil.com으로 간다(%2f·탭도 같은 갈래).
_HOST_RE = re.compile(r"[a-z0-9.-]+")
_GOODS_MAX = 1000
memo_collections = Table(
    "memo_collections",
    _METADATA,
    Column("id", Integer, primary_key=True, autoincrement=True),
    Column("title", String(_TITLE_MAX), nullable=False, unique=True),  # 도구 인자·카드 제목 앞부분
    Column("description", Text, nullable=False),  # 언제 쓰는 자료인지(모델이 읽음)
    Column("source_url", String(_SOURCE_URL_MAX)),  # 관련 Yes24 페이지(선택)
    Column("period_start", Date),
    Column("period_end", Date),
    Column("active", Boolean, nullable=False, default=False),
    # 항상 참고 — 켜진 묶음이면 항목 전문을 매 턴 지시문에 싣는다(조회 판단을 거치지 않는다).
    Column("always_include", Boolean, nullable=False, default=False),
    Column("revision", Integer, nullable=False),  # 묶음과 그 안 메모의 모든 쓰기마다 +1
    Column("updated_by", String(64), nullable=False),  # 관리자 계정명 | "cli"
    Column("updated_at", DateTime, nullable=False),  # UTC naive
    Column("created_at", DateTime, nullable=False),
    mysql_engine="InnoDB",
    mysql_charset="utf8mb4",
)
memos = Table(
    "memos",
    _METADATA,
    Column("id", Integer, primary_key=True, autoincrement=True),  # 삽입 순서 = 표시 순서
    Column("collection_id", Integer, nullable=False),
    Column("content", Text, nullable=False),
    Column("about", String(_SHORT_MAX)),  # 누가(말한 사람/대상 인물)
    Column("link", String(_SOURCE_URL_MAX)),  # 관련 Yes24 페이지 하나(선택) — 카드 링크 1순위
    # ",123,456," 꼴(앞뒤 쉼표) 또는 '' — LIKE '%,{n},%'가 경계를 지킨다(12가 123에 걸리지 않게).
    Column("goods_nos", String(_GOODS_MAX), nullable=False, server_default=""),
    Column("updated_by", String(64), nullable=False),
    Column("updated_at", DateTime, nullable=False),
    Column("created_at", DateTime, nullable=False),
    Index("ix_memos_collection", "collection_id", "id"),
    mysql_engine="InnoDB",
    mysql_charset="utf8mb4",
)
COLLECTION_FIELDS = (
    "title", "description", "source_url", "period_start", "period_end", "active", "always_include",
)
MEMO_FIELDS = ("content", "about", "link", "goods_nos")


def mysql_ddl() -> str:
    """운영 적용 승인자가 볼 MySQL DDL(`init --print`)."""
    dialect = mysql.dialect()
    parts = [CreateTable(table).compile(dialect=dialect) for table in (memo_collections, memos)]
    parts += [CreateIndex(index).compile(dialect=dialect) for index in memos.indexes]
    return "\n".join(f"{str(part).strip()};" for part in parts)


# ── 검사 ───────────────────────────────────────────────────────────────────

# 개인 연락처 형태(D5). 값은 NFKC로 정규화한 뒤 본다(전각 숫자·＠·．). 공개 직함·약력은 걸리지
# 않는다 — ISBN(97x·89x)·날짜(20xx)·상품번호(0으로 시작하지 않음)·가격(쉼표)은 전화 패턴의
# "숫자 아닌 것 뒤의 0" 조건과 겹치지 않는다. 구분자는 하이픈·점·밑줄·닫는 괄호·대시류
# (U+2010~2015) 하나와 그 앞뒤 공백·탭(줄은 넘지 않는다)이고, 지역번호는 괄호로 감쌀 수 있다
# ("(02)1234-5678"·"02)1234-5678"). 국제 표기(+82)는 지역번호 앞 0을 빼거나 괄호로 감싼
# 꼴("+82 (0)10-…")도 받는다.
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_SEP = r"[ \t]*[-._)‐-―]?[ \t]*"
_AREA = r"(?:1[016789]|2|[3-6][1-5]|70)"
_PHONE_RE = re.compile(
    rf"(?:(?<![\d+])\(?0|\+82{_SEP}\(?0?\)?){_AREA}{_SEP}\d{{3,4}}{_SEP}\d{{4}}(?!\d)"
)
_GOODS_RE = re.compile(r"^[1-9][0-9]*$")  # 상품번호 = 0으로 시작하지 않는 ASCII 숫자
_GOODS_SPLIT_RE = re.compile(r"[\s,;/]+")
_BLANK_LINES_RE = re.compile(r"\n(?:[ \t]*\n)+")
_FORMATS = (".xlsx", ".csv")
_EXAMPLE_ROWS = 3  # 묶은 오류 문구에 예로 드는 행 번호 개수(표시용)
_PII_HINT = "이메일·전화번호 형태가 있습니다. 개인 연락처는 받지 않습니다"
_FIELD_LABELS = {"content": "내용", "about": "누가", "link": "링크", "goods_nos": "상품번호"}


def contains_pii(text: str) -> bool:
    """이메일·전화번호 형태 판정 — 메모·열 이름·시트 이름·묶음 메타가 같은 함수를 쓴다."""
    normalized = unicodedata.normalize("NFKC", text)
    return bool(_EMAIL_RE.search(normalized) or _PHONE_RE.search(normalized))


def _url_problem(url: str, settings: Settings) -> str | None:
    """메모 링크·묶음 관련 페이지 검사 — 카드 링크로 고객에게 가는 주소라 Yes24 주소만 받는다."""
    if not url:
        return None
    parts = urlsplit(url)
    if (
        len(url) > _SOURCE_URL_MAX
        or any(char.isspace() or ord(char) < 32 for char in url)  # urlsplit은 탭·개행을 조용히 지운다
        or parts.scheme not in ("https", "http")  # javascript://www.yes24.com/… 는 호스트만 보면 통과한다
        or parts.username is not None
        or not _HOST_RE.fullmatch(parts.hostname or "")
        or not is_allowed_host(parts.hostname, allowed_domain(settings.yes24_base_url))
    ):
        return "Yes24 주소여야 합니다"
    if "#" in url:
        return "주소에 '#' 조각은 넣지 않습니다"
    return None


@dataclass
class MemoIssue:
    """저장을 막는 문제 하나. 메시지에 값 원문은 싣지 않는다(위치만)."""

    rule: str
    message: str
    row: int | None = None  # 파일의 행 번호(1 기준) — 묶인 오류면 첫 행
    column: str | None = None
    count: int = 1  # 같은 규칙·같은 열에 걸린 행 수


@dataclass
class MemoDraft:
    content: str
    about: str | None = None
    link: str | None = None
    goods_nos: list[str] = field(default_factory=list)
    row: int | None = None  # 가져오기의 파일 행 번호(오류 안내용)

    def __post_init__(self) -> None:
        # 상품번호 정규화는 여기 한 번이다(앞뒤 공백·빈 토큰 제거) — 검사·저장은 이 값만 읽는다.
        self.goods_nos = [goods.strip() for goods in self.goods_nos if goods and goods.strip()]


def _goods_text(goods: list[str]) -> str:
    """상품번호 목록 → 저장 값(겹침 제거, 앞뒤 쉼표)."""
    return f",{','.join(dict.fromkeys(goods))}," if goods else ""


def _problems(draft: MemoDraft, settings: Settings) -> list[tuple[str, str, str]]:
    """(규칙, 필드, 문제) 목록 — 저장되는 값만 본다. 가져오기와 손 입력이 같은 규칙을 쓴다."""
    found = []
    content = (draft.content or "").strip()
    if not content:
        found.append(("empty", "content", "내용이 비었습니다"))
    elif len(content) > settings.memo_max_chars:
        found.append(("length", "content", f"너무 깁니다(최대 {settings.memo_max_chars}자)"))
    if len((draft.about or "").strip()) > _SHORT_MAX:
        found.append(("length", "about", f"너무 깁니다(최대 {_SHORT_MAX}자)"))
    for key in ("content", "about", "link"):
        if contains_pii(getattr(draft, key) or ""):
            found.append(("pii", key, _PII_HINT))
    if problem := _url_problem((draft.link or "").strip(), settings):
        found.append(("link", "link", problem))
    if not all(_GOODS_RE.match(goods) for goods in draft.goods_nos):
        problem = "상품번호가 아닌 값이 있습니다(0으로 시작하지 않는 숫자)"
        found.append(("goods", "goods_nos", problem))
    elif len(_goods_text(draft.goods_nos)) > _GOODS_MAX:
        found.append(("goods", "goods_nos", f"상품번호가 너무 많습니다(최대 {_GOODS_MAX}자)"))
    return found


def validate_memo(draft: MemoDraft, settings: Settings) -> list[MemoIssue]:
    """메모 한 건 검사(어드민 손 입력·편집, 저장 직전 재검사)."""
    return [
        MemoIssue(rule, f"{_FIELD_LABELS[key]}: {problem}.", column=key)
        for rule, key, problem in _problems(draft, settings)
    ]


def _memo_values(draft: MemoDraft) -> dict[str, Any]:
    """검사를 통과한 메모 → 저장 값."""
    return {
        "content": draft.content.strip(),
        "about": (draft.about or "").strip() or None,
        "link": (draft.link or "").strip() or None,
        "goods_nos": _goods_text(draft.goods_nos),
    }


# ── 가져오기 ───────────────────────────────────────────────────────────────


@dataclass
class ImportMapping:
    # 1개 이상. 파일 열 순서대로 "열이름: 값" 줄로 잇는다(빈 값 생략, 열이 하나면 값만)
    content_columns: list[str]
    about_column: str | None = None  # 누가
    goods_column: str | None = None  # 상품번호(칸 하나에 여러 개 가능)
    link_column: str | None = None  # 링크(Yes24 주소)


@dataclass
class ImportResult:
    file_name: str
    file_sha256: str  # 원본 바이트 해시(감사·표시용)
    sheets: list[str] = field(default_factory=list)  # 파일의 시트 전부(CSV는 [])
    sheet: str | None = None
    columns: list[str] = field(default_factory=list)  # 머리글(매핑 선택지)
    sample_rows: list[list[str]] = field(default_factory=list)  # 앞 N행 원시값
    memos: list[MemoDraft] = field(default_factory=list)  # mapping이 없으면 []
    skipped_rows: int = 0  # 매핑한 내용 열이 전부 빈 행
    notes: list[str] = field(default_factory=list)
    errors: list[MemoIssue] = field(default_factory=list)
    mapping: ImportMapping | None = None


def _cell_text(value: object) -> str:
    """셀 값을 문자열로 정규화한다."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if value.time() == time() else "%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # ISBN·상품번호가 지수 표기(9.79E+12)로 새지 않게
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return _BLANK_LINES_RE.sub("\n", text).strip()


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for name in names:
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return result


def _read_rows(
    raw_rows: Iterable[Sequence[object]], settings: Settings, issues: list[MemoIssue]
) -> tuple[list[str], list[tuple[int, list[str]]]] | None:
    """행 반복자 → (열 이름, [(파일 행 번호, 값)]). 담는 칸은 (행 상한+1)×열 상한까지다.

    머리글은 첫 비어 있지 않은 행이다. 빈 행은 담지 않고, 데이터 행이 상한을 넘으면 그 자리에서
    멈춘다(멀리 떨어진 칸 하나가 격자 전체를 만들게 하지 않는다).
    """
    max_rows, max_cols = settings.memo_import_max_rows, settings.memo_import_max_columns
    header: list[str] | None = None
    data: list[tuple[int, list[str]]] = []
    for number, raw in enumerate(raw_rows, start=1):
        if any(_cell_text(value) for value in raw[max_cols:]):
            issues.append(MemoIssue(
                "max_columns",
                f"{number}행에 {max_cols}번째 열 밖의 값이 있습니다. 표는 {max_cols}열까지 "
                "받습니다 — 표 밖 먼 칸에 남은 값이 없는지 확인해 주세요.",
                number,
            ))
            return None
        values = [_cell_text(value) for value in raw[:max_cols]]
        if not any(values):
            continue
        if header is None:
            header = values
            continue
        data.append((number, values))
        if len(data) > max_rows:
            issues.append(MemoIssue("max_rows", f"데이터 행이 상한({max_rows})을 넘습니다."))
            return None
    if header is None:
        issues.append(MemoIssue("empty", "머리글 행이 없습니다."))
        return None

    width = max(len(header), *(len(values) for _number, values in data), 0)
    header = header + [""] * (width - len(header))
    data = [(number, values + [""] * (width - len(values))) for number, values in data]
    keep, names = [], []
    for col in range(width):
        label = header[col]
        if not label:
            if not any(values[col] for _number, values in data):
                continue  # 머리글도 값도 없는 열은 버린다
            label = f"열{col + 1}"
        keep.append(col)
        names.append(label)
    return _dedupe(names), [(number, [values[col] for col in keep]) for number, values in data]


def _xlsx_rows(
    data: bytes, wanted: str | None, result: ImportResult, settings: Settings,
    issues: list[MemoIssue],
) -> Iterator[Sequence[object]] | None:
    # 도구 import 경로(toolsets → memos)가 openpyxl을 싣지 않게 여기서 불러온다.
    import openpyxl

    try:
        # read_only = 시트 셀 스트리밍(실제 셀만 읽는다). 공유 문자열(sharedStrings)·스타일은
        # load_workbook이 통째로 올리므로, 열기 전에 압축 해제 크기 합을 상한과 견준다(zip이 적은
        # 크기를 넘겨 풀리지 않는다 — CRC·크기 검사). XML 폭탄은 defusedxml(의존성)이 막는다.
        # 열은 상한+1까지만 요청한다 — 상한 바로 다음 열에 값이 있으면 표가 넓다고 알리고
        # (_read_rows), 그보다 먼 외딴 칸은 읽지 않는다. 수식 칸은 엑셀이 저장한 계산값만 읽는다.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unzipped = sum(info.file_size for info in archive.infolist())
        if unzipped > settings.memo_import_max_unzipped_bytes:
            limit_mb = settings.memo_import_max_unzipped_bytes / 1_000_000
            issues.append(MemoIssue("size", f"압축을 푼 크기가 너무 큽니다(최대 {limit_mb:g} MB)."))
            return None
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — 손상·비엑셀 파일은 형식 오류 하나로 안내한다
        issues.append(MemoIssue("format", "엑셀(xlsx) 파일을 열 수 없습니다."))
        return None
    result.sheets = list(workbook.sheetnames)
    name = wanted or result.sheets[0]
    if name not in result.sheets:
        issues.append(MemoIssue("unknown_sheet", "없는 시트입니다."))
        return None
    result.sheet = name
    return workbook[name].iter_rows(
        max_col=settings.memo_import_max_columns + 1, values_only=True
    )


def _csv_rows(data: bytes, issues: list[MemoIssue]) -> Iterator[Sequence[object]] | None:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        issues.append(MemoIssue(
            "encoding",
            "CSV는 UTF-8만 받습니다. 엑셀에서 '다른 이름으로 저장 → CSV UTF-8'로 저장하거나 "
            "xlsx 그대로 올려 주세요.",
        ))
        return None
    return csv.reader(io.StringIO(text))


def _grouped(rule: str, label: str, numbers: list[int], problem: str) -> MemoIssue:
    """같은 규칙·같은 열을 한 줄로 — 행 번호는 앞 몇 개만 예로 들고 값은 싣지 않는다."""
    example = "·".join(str(number) for number in numbers[:_EXAMPLE_ROWS])
    more = " 등" if len(numbers) > _EXAMPLE_ROWS else ""
    return MemoIssue(
        rule,
        f"'{label}' {len(numbers)}개 행({example}행{more})에 {problem} — 매핑에서 빼거나 값을 "
        "고쳐 다시 올려 주세요.",
        numbers[0], label, len(numbers),
    )


def read_import(
    file_name: str, data: bytes, *, sheet: str | None, mapping: ImportMapping | None,
    settings: Settings,
) -> ImportResult:
    """엑셀/CSV → 메모 초안. `errors`가 비어 있어야 저장할 수 있다.

    mapping이 없으면 시트·열·원시 샘플까지만 읽는다(매핑을 고르는 미리보기 1단계).
    """
    result = ImportResult(file_name=file_name, file_sha256=hashlib.sha256(data).hexdigest())
    issues = result.errors
    if len(data) > settings.memo_import_max_bytes:
        limit_mb = settings.memo_import_max_bytes / 1_000_000
        issues.append(MemoIssue("size", f"파일이 너무 큽니다(최대 {limit_mb:g} MB)."))
        return result
    suffix = PurePath(file_name).suffix.lower()
    if suffix not in _FORMATS:
        issues.append(MemoIssue("format", "xlsx 또는 CSV(UTF-8) 파일만 받습니다."))
        return result
    raw = (
        _xlsx_rows(data, sheet, result, settings, issues)
        if suffix == ".xlsx" else _csv_rows(data, issues)
    )
    if raw is None:
        return result
    try:
        read = _read_rows(raw, settings, issues)
    except Exception as exc:  # noqa: BLE001 — 시트 XML은 읽는 도중에 파싱된다(read_only 지연 파싱)
        # 잘린 시트 XML·zip CRC 오류·CSV 칸 상한 초과 등 — 열 때와 같은 형식 오류(값 미포함).
        logger.warning(f"memo import unreadable file={file_name!r}: {type(exc).__name__}")
        message = "파일을 끝까지 읽을 수 없습니다(손상, 또는 CSV 칸이 너무 길거나 따옴표 어긋남)."
        issues.append(MemoIssue("format", message))
        return result
    if read is None:
        return result
    result.columns, rows = read
    result.sample_rows = [values for _number, values in rows[: settings.admin_page_size]]
    if result.sheet and contains_pii(result.sheet):
        issues.append(MemoIssue("pii", f"시트 이름에 {_PII_HINT}."))
    if not rows:
        issues.append(MemoIssue("empty", "표에 데이터 행이 없습니다."))
    elif mapping is not None:
        result.mapping = mapping
        _map_rows(result, rows, mapping, settings)
    return result


def _map_rows(
    result: ImportResult, rows: list[tuple[int, list[str]]], mapping: ImportMapping,
    settings: Settings,
) -> None:
    columns, issues = result.columns, result.errors
    named = [*mapping.content_columns, mapping.about_column, mapping.goods_column,
             mapping.link_column]
    unknown = [name for name in named if name is not None and name not in columns]
    if not mapping.content_columns or unknown:
        message = f"없는 열입니다: {', '.join(unknown)}" if unknown else "내용 열을 고르세요."
        issues.append(MemoIssue("unknown_column", message))
        return
    for name in dict.fromkeys(name for name in named if name):
        if contains_pii(name):
            position = columns.index(name) + 1
            issues.append(MemoIssue("pii", f"{position}번째 열 이름에 {_PII_HINT}."))
    content = [index for index, name in enumerate(columns) if name in mapping.content_columns]

    def cell(values: list[str], name: str | None) -> str:
        return values[columns.index(name)] if name else ""

    labels = {
        "content": "내용", "about": mapping.about_column, "link": mapping.link_column,
        "goods_nos": mapping.goods_column,
    }
    found: dict[tuple[str, str], tuple[str, list[int]]] = {}
    for number, values in rows:
        lines = [values[index] if len(content) == 1 else f"{columns[index]}: {values[index]}"
                 for index in content if values[index]]
        if not lines:
            result.skipped_rows += 1  # 건너뛴 행의 누가·링크·상품번호는 검사하지 않는다
            continue
        draft = MemoDraft(
            "\n".join(lines), cell(values, mapping.about_column) or None,
            cell(values, mapping.link_column) or None,
            _GOODS_SPLIT_RE.split(cell(values, mapping.goods_column)), number,
        )
        for rule, key, problem in _problems(draft, settings):
            found.setdefault((rule, labels[key]), (problem, []))[1].append(number)
        result.memos.append(draft)
    issues.extend(_grouped(rule, label, numbers, problem)
                  for (rule, label), (problem, numbers) in found.items())
    if result.skipped_rows:
        result.notes.append(f"내용이 빈 행 {result.skipped_rows}개를 건너뛰었습니다.")
    if not result.memos:
        message = "가져올 자료가 없습니다(고른 내용 열이 전부 비었습니다)."
        issues.append(MemoIssue("empty", message))


# ── 묶음 메타 ──────────────────────────────────────────────────────────────


@dataclass
class CollectionMeta:
    title: str
    description: str
    source_url: str | None = None
    period_start: date | None = None
    period_end: date | None = None


def validate_collection(meta: CollectionMeta, settings: Settings) -> list[MemoIssue]:
    issues: list[MemoIssue] = []

    def bad(key: str, message: str) -> None:
        issues.append(MemoIssue("meta", message, column=key))

    title = (meta.title or "").strip()
    if not title or len(title) > _TITLE_MAX:
        bad("title", f"이름은 1~{_TITLE_MAX}자입니다.")
    description = (meta.description or "").strip()
    description_max = settings.memo_description_max_chars
    if not description or len(description) > description_max:
        bad("description", f"설명은 1~{description_max}자입니다.")
    # 이름·설명은 지시문 목록으로, 관련 페이지는 메모 출처 url 바탕으로 모델·공개 출처에 그대로
    # 간다 — 메모와 같은 개인정보 검사를 한다.
    url = meta.source_url or ""
    for key, value in (("title", title), ("description", description), ("source_url", url)):
        if contains_pii(value):
            bad(key, f"{_PII_HINT}.")
    if problem := _url_problem(url, settings):
        bad("source_url", f"관련 페이지: {problem}.")
    if meta.period_start and meta.period_end and meta.period_start > meta.period_end:
        bad("period_end", "기간 끝이 시작보다 앞섭니다.")
    return issues


def _reject(issues: list[MemoIssue]) -> None:
    if issues:
        raise HTTPException(status_code=422, detail=[asdict(issue) for issue in issues])


# ── 저장 ───────────────────────────────────────────────────────────────────

# (action, 묶음 id, before, after) → (감사 INSERT 문장, 바인딩)
AuditHook = Callable[[str, int, dict | None, dict | None], tuple[str, tuple]]


def _utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _kst(value: datetime | None) -> str | None:
    if value is None:
        return None
    return value.replace(tzinfo=timezone.utc).astimezone(KST).strftime("%Y-%m-%d %H:%M")


def _iso(value: date | str | None) -> str | None:
    return value.isoformat() if isinstance(value, date) else value


def _period(start: Any, end: Any) -> dict | None:
    return None if start is None and end is None else {"start": _iso(start), "end": _iso(end)}


def _row(row: Any) -> dict[str, Any]:
    """DB 행 → JSON 직렬화 가능한 dict(날짜는 ISO, 시각은 KST 표시 문자열)."""
    data = dict(row._mapping)
    for key in ("updated_at", "created_at"):
        data[key] = _kst(data[key])
    for key in ("period_start", "period_end"):
        data[key] = _iso(data[key])
    if "active" in data:  # 묶음
        data["active"], data["always_include"] = bool(data["active"]), bool(data["always_include"])
    else:  # 메모 — 묶음 기간은 투영이 쓰는 한 덩어리로, 상품번호는 목록으로
        data["period"] = _period(data.pop("period_start"), data.pop("period_end"))
        data["goods_nos"] = [goods for goods in data["goods_nos"].split(",") if goods]
    return data


_MEMO_COUNT = (
    select(func.count()).where(memos.c.collection_id == memo_collections.c.id).scalar_subquery()
)
_COLLECTION_SELECT = select(memo_collections, _MEMO_COUNT.label("memo_count"))
_MEMO_SELECT = select(
    memos, memo_collections.c.title.label("collection"),
    memo_collections.c.source_url.label("collection_url"),
    memo_collections.c.period_start, memo_collections.c.period_end,
).join(memo_collections, memos.c.collection_id == memo_collections.c.id)
_ACTIVE = memo_collections.c.active.is_(True)


def _terms(query: str) -> list[str]:
    return [term for term in (_norm(word) for word in query.split()) if term]


def _matching(found: list[dict], query: str) -> list[dict]:
    """낱말마다 내용·누가에 들거나 상품번호와 정확히 같은 메모(NFKC·대소문자·공백 무시)."""
    terms = _terms(query)

    def hit(memo: dict) -> bool:
        text = "\x1f".join(_norm(memo[key] or "") for key in ("content", "about"))
        return all(term in text or term in memo["goods_nos"] for term in terms)

    return [memo for memo in found if hit(memo)]


def _meta_of(collection: dict, changes: dict) -> CollectionMeta:
    merged = {**collection, **changes}
    meta = CollectionMeta(**{item.name: merged[item.name] for item in fields(CollectionMeta)})
    for key in ("period_start", "period_end"):
        if isinstance(getattr(meta, key), str):
            setattr(meta, key, date.fromisoformat(getattr(meta, key)))
    return meta


def _clean_meta(values: dict[str, Any]) -> dict[str, Any]:
    """묶음 메타 → 저장 값(이름·설명 앞뒤 공백 제거, 빈 관련 페이지 = NULL)."""
    cleaned = {key: value.strip() if key in ("title", "description") else value
               for key, value in values.items()}
    if "source_url" in cleaned:
        cleaned["source_url"] = cleaned["source_url"] or None
    return cleaned


def _memo_audit(memo: dict, keys: Iterable[str]) -> dict:
    """메모 감사 — 버전이 없어 수정 이력을 추적할 길이 감사뿐이라 값을 남긴다(PII 검사 통과값)."""
    return {"memo_id": memo["id"], **{key: memo[key] for key in keys}}


def _check_fields(changes: dict, allowed: tuple[str, ...]) -> None:
    unknown = sorted(set(changes) - set(allowed))
    if unknown or not changes:
        raise HTTPException(status_code=422, detail=f"바꿀 수 없는 항목: {unknown}")


class MemoService:
    """메모 저장소(sqlite·MySQL 공용). 엔진은 처음 쓸 때 만든다."""

    def __init__(self, db_url: str, *, pool_size: int) -> None:
        self._db_url = db_url
        self._pool_size = pool_size
        self._engine: AsyncEngine | None = None

    @property
    def engine(self) -> AsyncEngine:
        if self._engine is None:
            options: dict[str, Any] = {"pool_size": self._pool_size, "pool_pre_ping": True}
            # MySQL은 RC로 연다 — RR의 갭 잠금이 서로 다른 묶음의 동시 교체를 교착시킨 이력이
            # 있다. 같은 묶음의 쓰기는 묶음 행의 조건부 UPDATE가 한 줄로 세운다.
            if make_url(self._db_url).get_backend_name() == "mysql":
                options["isolation_level"] = "READ COMMITTED"
                # 접속 수립 상한(session_service.mysql_pool_kwargs와 같은 값) — 드라이버 기본은
                # None(OS TCP 타임아웃)이라 RDS가 응답하지 않으면 실패 판정이 분 단위로 늘어진다.
                timeout = get_settings().mysql_connect_timeout_s
                options["connect_args"] = {"connect_timeout": timeout}
            self._engine = create_async_engine(self._db_url, **options)
        return self._engine

    async def close(self) -> None:
        if self._engine is not None:
            await self._engine.dispose()
            self._engine = None

    async def init_schema(self) -> None:
        async with self.engine.begin() as conn:
            await conn.run_sync(_METADATA.create_all, checkfirst=True)

    # ── 읽기 ──

    async def _read(self, stmt) -> list[dict]:
        async with self.engine.connect() as conn:
            return [_row(row) for row in (await conn.execute(stmt)).all()]

    async def list_collections(self, active_only: bool = False) -> list[dict]:
        stmt = _COLLECTION_SELECT.order_by(memo_collections.c.id)
        return await self._read(stmt.where(_ACTIVE) if active_only else stmt)

    async def list_memos(self, cid: int, query: str = "") -> dict:
        """어드민용 — 꺼진 묶음도 본다. 도구와 같은 일치 함수, 메모 id 순."""
        async with self.engine.connect() as conn:
            collection = await self._collection(conn, cid)
        found = await self._read(
            _MEMO_SELECT.where(memos.c.collection_id == cid).order_by(memos.c.id)
        )
        return {"collection": collection, "memos": _matching(found, query)}

    async def search(self, collection: str, query: str) -> dict:
        """활성 묶음 조회 — {"collections": 활성 묶음, "memos": 일치 메모 | None(모르는 묶음)}.

        묶음은 이름 부분 일치(여러 개 가능, 비우면 전부), 메모는 낱말 전부 포함.
        """
        active = await self.list_collections(active_only=True)
        chosen = [item["id"] for item in active if _norm(collection) in _norm(item["title"])]
        if collection.strip() and not chosen:
            return {"collections": active, "memos": None}
        found = await self._read(_MEMO_SELECT.where(memos.c.collection_id.in_(chosen)).order_by(
            memos.c.collection_id, memos.c.id
        ))
        return {"collections": active, "memos": _matching(found, query)}

    # ── 쓰기 ──

    @contextlib.asynccontextmanager
    async def _writing(self, cid: int, expected_revision: int, by: str):
        """모든 쓰기의 트랜잭션 — 첫 문장은 묶음 행의 조건부 UPDATE(revision+1)다.

        이 문장이 묶음 행에 X 잠금을 걸어 같은 묶음의 쓰기를 한 줄로 세운다(통째 교체의 DELETE가
        다른 쓰기와 엇갈리지 않는다). 바뀐 행이 0이면 없음(404)과 경쟁(409)을 가른다.
        """
        async with self.engine.begin() as conn:
            result = await conn.execute(
                update(memo_collections)
                .where(memo_collections.c.id == cid,
                       memo_collections.c.revision == expected_revision)
                .values(revision=memo_collections.c.revision + 1, updated_by=by,
                        updated_at=_utcnow())
            )
            if result.rowcount == 0:
                raise EditConflict(await self._collection(conn, cid))
            yield conn

    async def create_collection(
        self, meta: CollectionMeta, by: str, audit: AuditHook | None = None
    ) -> dict:
        """꺼진 상태·revision 1로 만든다. 이름이 겹치면 409."""
        _reject(validate_collection(meta, get_settings()))
        now = _utcnow()
        values = {**_clean_meta(asdict(meta)), "active": False, "always_include": False, "revision": 1,
                  "updated_by": by, "updated_at": now, "created_at": now}
        async with self.engine.begin() as conn:
            try:
                result = await conn.execute(insert(memo_collections).values(**values))
            except IntegrityError:
                raise HTTPException(status_code=409, detail="이미 있는 묶음 이름입니다.") from None
            created = await self._collection(conn, result.inserted_primary_key[0])
            await _audit(conn, audit, "create", created["id"], None,
                         {key: created[key] for key in COLLECTION_FIELDS})
        return created

    async def update_collection(
        self, cid: int, changes: dict[str, Any], expected_revision: int, by: str,
        audit: AuditHook | None = None,
    ) -> dict:
        """메타 편집·켜기/끄기. 바뀐 메타로 규칙을 다시 검사한다(끄기만은 검사 없이 통과)."""
        _check_fields(changes, COLLECTION_FIELDS)
        async with self._writing(cid, expected_revision, by) as conn:
            before = await self._collection(conn, cid)
            # 끄기는 검사하지 않는다 — 사고 대응 수단이라, 설정이 바뀌어 저장된 메타가 지금
            # 규칙에 안 맞아도 막히면 안 된다. 켜기·메타 편집만 다시 검사한다.
            if changes != {"active": False} and set(changes) != {"always_include"}:
                _reject(validate_collection(_meta_of(before, changes), get_settings()))
            try:
                await conn.execute(update(memo_collections).where(memo_collections.c.id == cid)
                                   .values(**_clean_meta(changes)))
            except IntegrityError:
                raise HTTPException(status_code=409, detail="이미 있는 묶음 이름입니다.") from None
            after = await self._collection(conn, cid)
            action = "update"
            if set(changes) == {"active"}:
                action = "activate" if changes["active"] else "deactivate"
            elif set(changes) == {"always_include"}:
                action = "pin" if changes["always_include"] else "unpin"
            await _audit(conn, audit, action, cid, {key: before[key] for key in changes},
                         {key: after[key] for key in changes})
            return after

    async def delete_collection(
        self, cid: int, expected_revision: int, by: str, audit: AuditHook | None = None
    ) -> None:
        """묶음과 그 메모를 지운다(되돌릴 수 없음 — owner·CLI --yes 전용)."""
        async with self._writing(cid, expected_revision, by) as conn:
            before = await self._collection(conn, cid)
            await conn.execute(delete(memos).where(memos.c.collection_id == cid))
            await conn.execute(delete(memo_collections).where(memo_collections.c.id == cid))
            await _audit(conn, audit, "delete", cid,
                         {"title": before["title"], "memo_count": before["memo_count"]}, None)

    async def create_memo(
        self, cid: int, draft: MemoDraft, expected_revision: int, by: str,
        audit: AuditHook | None = None,
    ) -> dict:
        _reject(validate_memo(draft, get_settings()))
        async with self._writing(cid, expected_revision, by) as conn:
            now = _utcnow()
            result = await conn.execute(insert(memos).values(
                collection_id=cid, **_memo_values(draft), updated_by=by, updated_at=now,
                created_at=now,
            ))
            memo = await self._memo(conn, cid, result.inserted_primary_key[0])
            await _audit(conn, audit, "memo_create", cid, None, _memo_audit(memo, MEMO_FIELDS))
            return {"collection": await self._collection(conn, cid), "memo": memo}

    async def update_memo(
        self, cid: int, memo_id: int, changes: dict[str, Any], expected_revision: int, by: str,
        audit: AuditHook | None = None,
    ) -> dict:
        _check_fields(changes, MEMO_FIELDS)
        async with self._writing(cid, expected_revision, by) as conn:
            before = await self._memo(conn, cid, memo_id)
            merged = {**before, **changes}
            draft = MemoDraft(merged["content"], merged["about"], merged["link"],
                              list(merged["goods_nos"]))
            _reject(validate_memo(draft, get_settings()))
            await conn.execute(update(memos).where(memos.c.id == memo_id).values(
                **_memo_values(draft), updated_by=by, updated_at=_utcnow(),
            ))
            after = await self._memo(conn, cid, memo_id)
            changed = [key for key in MEMO_FIELDS if before[key] != after[key]]
            await _audit(conn, audit, "memo_update", cid, _memo_audit(before, changed),
                         _memo_audit(after, changed))
            return {"collection": await self._collection(conn, cid), "memo": after}

    async def delete_memo(
        self, cid: int, memo_id: int, expected_revision: int, by: str,
        audit: AuditHook | None = None,
    ) -> dict:
        async with self._writing(cid, expected_revision, by) as conn:
            before = await self._memo(conn, cid, memo_id)
            await conn.execute(delete(memos).where(memos.c.id == memo_id))
            await _audit(conn, audit, "memo_delete", cid, _memo_audit(before, MEMO_FIELDS), None)
            return {"collection": await self._collection(conn, cid)}

    async def import_memos(
        self, cid: int, result: ImportResult, mode: str, expected_revision: int, by: str,
        audit: AuditHook | None = None,
    ) -> dict:
        """가져오기 — append는 뒤에 붙이고, replace는 옛 메모를 같은 트랜잭션에서 지우고 넣는다.

        켜진 묶음이면 저장 즉시 답변에 쓰인다(확인은 화면·CLI가 받는다). 감사엔 개수·파일 해시·
        매핑만 남기고 메모 값은 남기지 않는다. 초안은 저장 직전에 다시 검사한다.
        """
        if mode not in ("append", "replace"):
            raise HTTPException(status_code=422, detail="mode는 append 또는 replace입니다.")
        settings = get_settings()
        _reject(result.errors or [issue for draft in result.memos
                                  for issue in validate_memo(draft, settings)])
        if not result.memos:
            _reject([MemoIssue("empty", "가져올 자료가 없습니다.")])
        async with self._writing(cid, expected_revision, by) as conn:
            deleted = 0
            if mode == "replace":
                deleted = (await conn.execute(
                    delete(memos).where(memos.c.collection_id == cid)
                )).rowcount
            now = _utcnow()
            await conn.execute(insert(memos), [
                {"collection_id": cid, **_memo_values(draft), "updated_by": by,
                 "updated_at": now, "created_at": now}
                for draft in result.memos
            ])
            inserted = len(result.memos)
            mapping = asdict(result.mapping) if result.mapping else None
            await _audit(conn, audit, f"import_{mode}", cid, None, {
                "file_name": result.file_name, "file_sha256": result.file_sha256,
                "sheet": result.sheet, "mapping": mapping, "inserted": inserted,
                "deleted": deleted,
            })
            return {"collection": await self._collection(conn, cid), "inserted": inserted,
                    "deleted": deleted}

    @staticmethod
    async def _collection(conn, cid: int) -> dict:
        row = (await conn.execute(
            _COLLECTION_SELECT.where(memo_collections.c.id == cid)
        )).first()
        if row is None:
            raise HTTPException(status_code=404, detail="없는 묶음입니다.")
        return _row(row)

    @staticmethod
    async def _memo(conn, cid: int, memo_id: int) -> dict:
        row = (await conn.execute(
            _MEMO_SELECT.where(memos.c.id == memo_id, memos.c.collection_id == cid)
        )).first()
        if row is None:
            raise HTTPException(status_code=404, detail="이 묶음에 없는 자료입니다.")
        return _row(row)


async def _audit(conn, audit: AuditHook | None, action: str, cid: int, before, after) -> None:
    """감사는 데이터 쓰기와 같은 연결·트랜잭션이다 — 실패하면 쓰기도 롤백된다."""
    if audit is not None:
        statement, params = audit(action, cid, before, after)
        await conn.exec_driver_sql(statement, params)


_service: MemoService | None = None


def get_memo_service() -> MemoService:
    global _service
    if _service is None:
        settings = get_settings()
        _service = MemoService(settings.session_db_url, pool_size=settings.memo_pool_size)
    return _service


async def close_memo_service() -> None:
    """lifespan 종료 훅(toolsets.ACLOSE_HOOKS) — 엔진을 만든 적이 없으면 no-op."""
    global _service
    if _service is not None:
        await _service.close()
        _service = None


# ── 투영(도구·CLI 공용) ───────────────────────────────────────────────

_TRUNCATED = "결과가 길어 앞쪽 항목만 담았습니다 — query로 좁혀 다시 부르세요."
_NO_MATCH = "맞는 항목이 없습니다 — 낱말을 줄이거나 collection 없이 다시 찾으세요."


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text).casefold())


def memo_url(memo: dict, settings: Settings) -> str:
    """메모 출처 url — 메모 id가 정체성이다(편집 = 같은 출처 갱신, 통째 교체 = 새 출처).

    바탕(카드 링크가 가는 곳)은 링크 → 묶음 관련 페이지 → 첫 상품 페이지 → Yes24 홈 순이다.
    조각(#notice-id)이 상품 출처와 url을 갈라 같은 턴 상품 상세가 메모로 덮이지 않는다.
    """
    goods = memo.get("goods_nos") or []
    base = (
        memo.get("link") or memo.get("collection_url")
        or (product_url(settings.yes24_base_url, goods[0]) if goods
            else settings.yes24_base_url.rstrip("/"))
    )
    return f"{base}#notice-{memo['id']}"


def memo_view(memo: dict, settings: Settings) -> dict:
    """메모 → 출처 항목(등록 전). 조회 도구와 CLI가 같은 투영을 쓴다. 공개 제목 = 묶음 이름."""
    view: dict[str, Any] = {
        "type": MEMO_SOURCE_TYPE,
        "fidelity": DETAIL_FIDELITY,
        "title": memo["collection"],
        "url": memo_url(memo, settings),
        "snippet": memo["content"],
    }
    view.update({key: memo[key] for key in ("about", "period") if memo.get(key)})
    view["checked_at"] = memo["updated_at"]
    return view


def search_view(found: list[dict], settings: Settings) -> dict:
    """조회 도구 응답(출처 등록 전) — 앞에서부터 글자 예산까지, 첫 메모가 넘치면 잘라 담는다."""
    budget = settings.memo_result_max_chars
    items: list[dict] = []
    used, clipped = 0, False
    for memo in found:
        item = memo_view(memo, settings)
        if memo["goods_nos"]:
            item["goods_urls"] = [product_url(settings.yes24_base_url, goods)
                                  for goods in memo["goods_nos"]]
        if used + len(item["snippet"]) > budget:
            if items:
                break
            item["snippet"], clipped = item["snippet"][:budget] + "…", True
        items.append(item)
        used += len(item["snippet"])
    result = {
        "status": "ok",
        "items": items,
        "total_matches": len(found),
        "returned": len(items),
        "truncated": clipped or len(items) < len(found),
    }
    if result["truncated"]:
        result["message"] = _TRUNCATED
    elif not found:
        result["message"] = _NO_MATCH
    return result


def catalog_lines(collections: list[dict]) -> list[str]:
    """지시문용 묶음 한 줄 = 이름 — 설명 첫 문장(첫 줄의 ' / ' 앞, 마침표에서 끊음).

    지시문은 매 라운드 재전송되는 비용이라 설명 전문은 싣지 않는다(전문은 어드민·CLI용).
    """
    lines = []
    for item in collections:
        head = (item["description"] or "").strip().split("\n")[0].split(" / ")[0]
        first = re.split(r"(?<=[.])\s", head, maxsplit=1)[0]
        lines.append(f"{item['title']} — {first}" if first else item["title"])
    return lines
